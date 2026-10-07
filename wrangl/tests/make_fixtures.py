import os, io, math, random, zipfile, wave, struct, textwrap
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4, letter, landscape
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfform
from PIL import Image, ImageDraw, ImageFont, ImageFilter
random.seed(7)
D = os.path.join(os.path.dirname(__file__), 'fixtures'); os.makedirs(D, exist_ok=True)
P = lambda n: os.path.join(D, n)
LOREM = ("Lorem ipsum dolor sit amet, consectetur adipiscing elit, sed do eiusmod tempor incididunt ut labore et dolore magna aliqua. "
         "Ut enim ad minim veniam, quis nostrud exercitation ullamco laboris nisi ut aliquip ex ea commodo consequat. "
         "Duis aute irure dolor in reprehenderit in voluptate velit esse cillum dolore eu fugiat nulla pariatur.")
def para(c, text, x, y, w=480, size=11, lead=15, font='Helvetica'):
    c.setFont(font, size)
    for line in textwrap.wrap(text, int(w/(size*0.5))):
        c.drawString(x, y, line); y -= lead
    return y

# 1. report10.pdf with outline
c = canvas.Canvas(P('report10.pdf'), pagesize=A4); c.setTitle('Quarterly Report'); c.setAuthor('Test Author')
chap = {1:'Chapter 1: Introduction', 4:'Chapter 2: Methods', 7:'Chapter 3: Results'}
for i in range(1, 11):
    if i in chap:
        key = f'ch{i}'; c.bookmarkPage(key); c.addOutlineEntry(chap[i], key, level=0)
    if i == 2: c.bookmarkPage('s2'); c.addOutlineEntry('Section 1.1: Background', 's2', level=1)
    c.setFont('Helvetica-Bold', 22); c.drawString(72, 780, chap.get(i, f'Report page {i}'))
    c.setFont('Helvetica', 9); c.setFillGray(.4); c.drawString(72, 40, f'Footer text page {i}'); c.setFillGray(0)
    y = para(c, LOREM, 72, 740)
    y = para(c, LOREM + ' Unique marker word: alpha%d.' % i, 72, y-10)
    c.setFont('Times-Italic', 12); c.drawString(72, y-20, 'An italic Times line with numbers 1,234.56 and email test@example.com on page %d.' % i)
    c.showPage()
c.save()

# 2. invoices.pdf
c = canvas.Canvas(P('invoices.pdf'), pagesize=letter)
for i in range(1, 7):
    if i % 2 == 1:
        c.setFont('Helvetica-Bold', 20); c.drawString(72, 720, f'INVOICE No. INV-{1000+i}'); c.setFont('Helvetica', 12); c.drawString(72, 690, f'Customer {chr(64+i)} Ltd'); c.drawString(72, 670, 'Total due: Rs 12,500')
    else:
        c.setFont('Helvetica', 12); c.drawString(72, 720, f'Continuation sheet for invoice {1000+i-1}'); c.drawString(72, 700, 'Line items continued...')
    c.showPage()
c.save()

# 3. a.pdf b.pdf
for nm, col in (('a', (0.8, 0.9, 1)), ('b', (1, 0.9, 0.8))):
    c = canvas.Canvas(P(nm + '.pdf'), pagesize=A4)
    for i in range(1, 4):
        c.setFillColorRGB(*col); c.rect(0, 0, 595, 842, fill=1, stroke=0); c.setFillColorRGB(0, 0, 0)
        c.setFont('Helvetica-Bold', 40); c.drawString(100, 700, f'{nm.upper()}{i}'); c.showPage()
    c.save()

# 4. spread.pdf (landscape 2-page spreads)
c = canvas.Canvas(P('spread.pdf'), pagesize=landscape(A4))
for i in range(1, 4):
    c.setFont('Helvetica-Bold', 30); c.drawString(120, 400, f'LEFT {i}'); c.drawString(560, 400, f'RIGHT {i}'); c.line(421, 0, 421, 595); c.showPage()
c.save()

# 5. form.pdf (AcroForm)
c = canvas.Canvas(P('form.pdf'), pagesize=A4)
c.setFont('Helvetica-Bold', 18); c.drawString(72, 780, 'Application Form')
c.setFont('Helvetica', 12); c.drawString(72, 740, 'Full name:'); c.drawString(72, 700, 'Email:'); c.drawString(72, 660, 'I agree to terms:')
form = c.acroForm
form.textfield(name='fullname', tooltip='Full name', x=180, y=730, width=250, height=22, borderStyle='inset', forceBorder=True)
form.textfield(name='email', tooltip='Email', x=180, y=690, width=250, height=22, borderStyle='inset', forceBorder=True)
form.checkbox(name='agree', tooltip='Agree', x=180, y=655, size=18, forceBorder=True)
c.showPage(); c.save()

# 6. images.pdf (big photos -> compressible)
def photo(w, h, seed=1):
    random.seed(seed)
    im = Image.new('RGB', (w, h))
    px = im.load()
    for y in range(h):
        for x in range(w):
            px[x, y] = (int(127+120*math.sin(x/37.0+seed)), int(127+120*math.sin(y/29.0)), int(127+120*math.sin((x+y)/53.0)))
    im = im.filter(ImageFilter.GaussianBlur(1))
    d = ImageDraw.Draw(im)
    for _ in range(60):
        x, y = random.randint(0, w), random.randint(0, h); r = random.randint(5, 60)
        d.ellipse([x-r, y-r, x+r, y+r], fill=tuple(random.randint(0, 255) for _ in range(3)))
    return im
c = canvas.Canvas(P('images.pdf'), pagesize=A4)
for i in range(3):
    im = photo(1600, 1100, i+1); buf = io.BytesIO(); im.save(buf, 'JPEG', quality=95); buf.seek(0)
    c.drawImage(ImageReader(buf), 40, 300, width=515, height=354); c.setFont('Helvetica', 12); c.drawString(40, 270, f'Photo {i+1}'); c.showPage()
c.save()
photo(1600, 1100, 5).save(P('photo.jpg'), quality=92)

# 7. scan.pdf (image of text, for OCR)
im = Image.new('L', (1240, 1754), 255); d = ImageDraw.Draw(im)
try:
    f = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 44); f2 = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf', 64)
except Exception:
    f = f2 = ImageFont.load_default()
d.text((120, 150), 'Scanned Document', fill=0, font=f2)
y = 300
for line in ['This page is only an image.', 'The quick brown fox jumps over the lazy dog.', 'Invoice number 48213 total 1,250.00', 'Contact: hello@example.com']:
    d.text((120, y), line, fill=0, font=f); y += 90
im = im.filter(ImageFilter.GaussianBlur(0.6)); im.save(P('scan.png'))
c = canvas.Canvas(P('scan.pdf'), pagesize=A4); c.drawImage(P('scan.png'), 0, 0, width=595.28, height=841.89); c.showPage(); c.save()

# images for image tools
im = Image.new('RGB', (800, 600), (245, 245, 245)); d = ImageDraw.Draw(im)
d.ellipse([250, 120, 550, 420], fill=(200, 40, 40)); d.rectangle([330, 300, 470, 520], fill=(30, 90, 200))
im.save(P('cutout.png'))
im2 = Image.new('RGB', (600, 400), (30, 120, 200)); d2 = ImageDraw.Draw(im2)
for k in range(0, 600, 20): d2.line([k, 0, k, 400], fill=(40, 130, 210), width=1)
d2.text((200, 180), 'WATERMARK', fill=(255, 255, 255), font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf', 48))
im2.save(P('watermarked.png'))
Image.new('RGB', (400, 300), (255, 255, 255)).save(P('white.png'))

# 8. docx
import docx
from docx.shared import Pt, Inches
dd = docx.Document(); dd.add_heading('Sample Document', 0); dd.add_paragraph('This is a paragraph with ').add_run('bold text').bold = True
p = dd.add_paragraph(); r = p.add_run('italic and underlined'); r.italic = True; r.underline = True
dd.add_heading('Section One', 1); dd.add_paragraph(LOREM)
dd.add_paragraph('First bullet', style='List Bullet'); dd.add_paragraph('Second bullet', style='List Bullet'); dd.add_paragraph('Numbered one', style='List Number'); dd.add_paragraph('Numbered two', style='List Number')
t = dd.add_table(rows=3, cols=3); t.style = 'Table Grid'
for i in range(3):
    for j in range(3): t.cell(i, j).text = f'R{i+1}C{j+1}'
dd.add_paragraph('Price: ₹1,250 and a rupee sign.')
dd.add_picture(P('cutout.png'), width=Inches(2.5))
for i in range(25): dd.add_paragraph(f'Paragraph {i+1}. ' + LOREM)
dd.save(P('sample.docx'))

# 9. xlsx
import openpyxl
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Sales'
ws.append(['Region', 'Q1', 'Q2', 'Q3', 'Total'])
for i, r in enumerate(['North', 'South', 'East', 'West', 'Central']):
    ws.append([r, 100+i*10, 120+i*7, 90+i*13, f'=SUM(B{i+2}:D{i+2})'])
ws2 = wb.create_sheet('Notes'); ws2.append(['Note', 'Value']); ws2.append(['Hello', 42])
wb.save(P('sample.xlsx'))
open(P('sample.csv'), 'w').write('Name,Qty,Price\nWidget,4,2.50\n"Gadget, large",10,19.99\nGizmo,1,100\n')

# 10. pptx
from pptx import Presentation
from pptx.util import Inches as I, Pt as Pp
pr = Presentation(); pr.slide_width = I(13.333); pr.slide_height = I(7.5)
s = pr.slides.add_slide(pr.slide_layouts[0]); s.shapes.title.text = 'Deck Title'; s.placeholders[1].text = 'A subtitle for the deck'
s = pr.slides.add_slide(pr.slide_layouts[1]); s.shapes.title.text = 'Bullets'; tf = s.placeholders[1].text_frame; tf.text = 'First point'; tf.add_paragraph().text = 'Second point'; tf.add_paragraph().text = 'Third point'
s = pr.slides.add_slide(pr.slide_layouts[5]); s.shapes.title.text = 'Picture slide'; s.shapes.add_picture(P('cutout.png'), I(3), I(2), width=I(4))
tb = s.shapes.add_textbox(I(1), I(6), I(6), I(1)); tb.text_frame.text = 'A custom text box'
pr.save(P('sample.pptx'))

# 11. md / html / txt
open(P('sample.md'), 'w').write('# Markdown Title\n\nSome **bold** and *italic* text with a [link](https://example.com).\n\n## List\n\n- one\n- two\n  - nested\n\n1. first\n2. second\n\n```\ncode block\n```\n\n| a | b |\n|---|---|\n| 1 | 2 |\n\n> A quote\n')
open(P('sample.html'), 'w').write('<html><head><style>h1{color:#c0392b} .box{background:#eef;border:1px solid #99c;padding:10px}</style></head><body><h1>HTML Title</h1><p>Paragraph with <b>bold</b> and <a href="https://example.com">a link</a>.</p><div class="box">Boxed content</div><ul><li>a</li><li>b</li></ul></body></html>')
open(P('sample.txt'), 'w').write('Plain text file.\nSecond line with unicode: café — ₹100.\n')

# 12. epub
def epub(path):
    z = zipfile.ZipFile(path, 'w')
    z.writestr('mimetype', 'application/epub+zip', compress_type=zipfile.ZIP_STORED)
    z.writestr('META-INF/container.xml', '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
    z.writestr('OEBPS/content.opf', '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>Test Book</dc:title><dc:creator>Jane Author</dc:creator><dc:identifier id="id">urn:uuid:1234</dc:identifier><dc:language>en</dc:language></metadata><manifest><item id="c1" href="ch1.xhtml" media-type="application/xhtml+xml"/><item id="c2" href="ch2.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="c1"/><itemref idref="c2"/></spine></package>')
    z.writestr('OEBPS/ch1.xhtml', '<?xml version="1.0"?><html xmlns="http://www.w3.org/1999/xhtml"><head><title>One</title></head><body><h1>Chapter One</h1><p>' + LOREM + '</p><p>Second paragraph in chapter one.</p></body></html>')
    z.writestr('OEBPS/ch2.xhtml', '<?xml version="1.0"?><html xmlns="http://www.w3.org/1999/xhtml"><head><title>Two</title></head><body><h1>Chapter Two</h1><p>' + LOREM + '</p></body></html>')
    z.close()
epub(P('sample.epub'))

# 13. audio (1.5 s tone, 16k mono)
w = wave.open(P('tone.wav'), 'wb'); w.setnchannels(1); w.setsampwidth(2); w.setframerate(16000)
for i in range(24000): w.writeframes(struct.pack('<h', int(8000*math.sin(2*math.pi*440*i/16000))))
w.close()

# 14. encrypted pdf via pypdf
from pypdf import PdfReader, PdfWriter
r = PdfReader(P('a.pdf')); wr = PdfWriter()
for pg in r.pages: wr.add_page(pg)
wr.encrypt('secret123', 'owner456', algorithm='AES-256'); wr.write(P('locked.pdf'))
wr = PdfWriter()
for pg in r.pages: wr.add_page(pg)
wr.encrypt('', 'owner456', use_128bit=True, permissions_flag=0x0004); wr.write(P('restricted.pdf'))
# damaged pdf
data = open(P('a.pdf'), 'rb').read(); open(P('broken.pdf'), 'wb').write(data[:len(data)//2] + b'\n%%EOF')
print('fixtures ready:', sorted(os.listdir(D)))
